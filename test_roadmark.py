"""
Test suite for the road-marking generator.

Network-free: exercises parsing, width logic, geometry, road selection and the
end-to-end pipeline (offline) so it runs anywhere.

Run:  python -m unittest test_roadmark -v
  or: python test_roadmark.py
"""
import math
import unittest
import xml.etree.ElementTree as ET
import numpy as np

import temp2updated as temp2


def _hline(highway, north_offset_m, way_id=1, **tags):
    """A short E-W way offset `north_offset_m` metres north of (18.0, 73.0)."""
    dlat = north_offset_m / 111_320.0
    t = {"highway": highway}; t.update(tags)
    return {"id": way_id, "tags": t,
            "geometry": [{"lon": 72.999, "lat": 18.0 + dlat},
                         {"lon": 73.001, "lat": 18.0 + dlat}]}


class TestParse(unittest.TestCase):
    NS = 'xmlns="http://www.opengis.net/kml/2.2"'

    def test_basic_point(self):
        kml = f'<kml {self.NS}><Placemark><name>A</name>' \
              f'<Point><coordinates>73.1,18.2,0</coordinates></Point></Placemark></kml>'
        r = temp2.parse_kml_string(kml)
        self.assertEqual(len(r), 1)
        self.assertAlmostEqual(r[0][0], 73.1)
        self.assertAlmostEqual(r[0][1], 18.2)
        self.assertEqual(r[0][2], "A")

    def test_single_quote_xmlns_and_prefixes(self):
        kml = ("<kml xmlns='http://www.opengis.net/kml/2.2' xmlns:gx='u'>"
               "<Document><gx:Tour/><Placemark>"
               "<Point><coordinates>1.5,2.5</coordinates></Point>"
               "</Placemark></Document></kml>")
        r = temp2.parse_kml_string(kml)
        self.assertEqual(len(r), 1)
        self.assertEqual(r[0][2], "Unnamed")     # no name -> "Unnamed"

    def test_multigeometry_and_missing_z(self):
        kml = (f'<kml {self.NS}><Placemark><name>M</name><MultiGeometry>'
               f'<Point><coordinates>1,2</coordinates></Point>'
               f'<Point><coordinates>3,4</coordinates></Point>'
               f'</MultiGeometry></Placemark></kml>')
        r = temp2.parse_kml_string(kml)
        self.assertEqual(len(r), 2)

    def test_ignores_non_point_geometry(self):
        kml = (f'<kml {self.NS}><Placemark><LineString>'
               f'<coordinates>1,2 3,4</coordinates></LineString></Placemark></kml>')
        self.assertEqual(temp2.parse_kml_string(kml), [])

    def test_malformed_raises(self):
        with self.assertRaises(ValueError):
            temp2.parse_kml_string("<kml><Placemark></kaml>")


class TestWidth(unittest.TestCase):
    def test_name_override_priority(self):
        w, src = temp2.estimate_width({"highway": "service"}, "Stop w=12", 1)
        self.assertEqual(w, 12.0)
        self.assertEqual(src, "name")

    def test_osm_width_tag(self):
        w, src = temp2.estimate_width({"width": "8.5"}, "x", 1)
        self.assertEqual(w, 8.5)
        self.assertEqual(src, "osm:width")

    def test_lanes_plus_shoulder(self):
        w, src = temp2.estimate_width({"highway": "trunk", "lanes": "2"}, "x", 1)
        self.assertAlmostEqual(w, 2 * temp2.LANE_WIDTH + temp2.SHOULDER_TOTAL["trunk"])
        self.assertTrue(src.startswith("lanes"))

    def test_class_default(self):
        w, src = temp2.estimate_width({"highway": "service"}, "x", 1)
        self.assertEqual(w, temp2.ROAD_WIDTH_DEFAULTS["service"])
        self.assertEqual(src, "default:service")

    def test_implausible_width_ignored(self):
        w, src = temp2.estimate_width({"highway": "residential", "width": "999"}, "x", 1)
        self.assertEqual(src, "default:residential")   # 999 rejected by sanity range


class TestGeometry(unittest.TestCase):
    def setUp(self):
        self.frame = temp2.LocalFrame(18.0, 73.0)
        self.straight = np.array([[-60.0, 0.0], [60.0, 0.0]])

    def _area(self, ring_xy):
        return temp2.polygon_area_m2([np.asarray(p) for p in ring_xy])

    def test_straight_rectangle_area(self):
        _, ring_xy, curved = temp2.build_polygon(self.frame, self.straight, 3.5, 3.5, 4.0)
        self.assertFalse(curved)
        self.assertAlmostEqual(self._area(ring_xy), 7.0 * 4.0, delta=0.5)   # width 7 x len 4

    def test_asymmetric_offsets(self):
        _, ring_xy, _ = temp2.build_polygon(self.frame, self.straight, 2.0, 5.0, 4.0)
        self.assertAlmostEqual(self._area(ring_xy), (2.0 + 5.0) * 4.0, delta=0.5)

    def test_curve_is_followed(self):
        R = 30.0
        arc = np.array([[R * math.sin(math.radians(a)), R * math.cos(math.radians(a)) - R]
                        for a in np.linspace(-40, 40, 17)])
        _, ring_xy, curved = temp2.build_polygon(self.frame, arc, 4.0, 4.0, 20.0)
        self.assertTrue(curved)
        # ribbon area ~= width * centreline length
        self.assertAlmostEqual(self._area(ring_xy), 8.0 * 20.0, delta=10)

    def test_hairpin_folds_to_rectangle(self):
        R = 2.5
        arc = np.array([[R * math.sin(math.radians(a)), R * math.cos(math.radians(a)) - R]
                        for a in np.linspace(-90, 90, 25)])
        _, ring_xy, curved = temp2.build_polygon(self.frame, arc, 3.5, 3.5, 8.0)
        self.assertEqual(len(ring_xy), 5)      # fell back to a clean rectangle
        self.assertFalse(curved)

    def test_polygon_centred_on_point(self):
        # point is the frame origin; polygon centroid should be ~ (0,0)
        _, ring_xy, _ = temp2.build_polygon(self.frame, self.straight, 3.0, 3.0, 4.0)
        pts = np.array(ring_xy[:-1])
        self.assertLess(abs(pts[:, 0].mean()), 0.5)
        self.assertLess(abs(pts[:, 1].mean()), 0.5)

    def test_bearing_north_and_east(self):
        self.assertAlmostEqual(temp2.road_bearing_at(np.array([[0, 0], [0, 100]])) % 180, 0, delta=1)
        self.assertAlmostEqual(temp2.road_bearing_at(np.array([[0, 0], [100, 0]])) % 180, 90, delta=1)


class TestChooseRoad(unittest.TestCase):
    def setUp(self):
        self.frame = temp2.LocalFrame(18.0, 73.0)

    def test_prefers_major_within_bias(self):
        els = [_hline("service", 3.0, 1), _hline("trunk", 4.0, 2)]
        el, _, _ = temp2.choose_road(self.frame, els)
        self.assertEqual(el["tags"]["highway"], "trunk")

    def test_picks_clearly_nearest(self):
        els = [_hline("service", 1.0, 1), _hline("trunk", 30.0, 2)]
        el, _, dist = temp2.choose_road(self.frame, els)
        self.assertEqual(el["tags"]["highway"], "service")
        self.assertAlmostEqual(dist, 1.0, delta=0.5)

    def test_no_driveable_returns_none(self):
        el, nodes, dist = temp2.choose_road(self.frame, [_hline("footway", 1.0, 1)])
        # footway is last-resort but still returned; a pure-empty list returns None
        el2, _, _ = temp2.choose_road(self.frame, [])
        self.assertIsNone(el2)


class TestPipeline(unittest.TestCase):
    def test_offline_produces_valid_kml(self):
        pts = [(73.0, 18.0, "A"), (73.0009, 18.0009, "B")]
        kml, boq = temp2.build_polygons(pts, offline=True, log=lambda *a: None)
        ET.fromstring(kml)                      # must be well-formed
        self.assertEqual(len(boq), 2)
        self.assertEqual(kml.count("<Polygon>"), 2)
        for b in boq:
            self.assertGreater(b["Area (m2)"], 0)

    def test_rate_flows_to_amount(self):
        pts = [(73.0, 18.0, "A")]
        _, boq = temp2.build_polygons(pts, offline=True, rate=100, log=lambda *a: None)
        self.assertAlmostEqual(boq[0]["Amount (Rs)"], boq[0]["Area (m2)"] * 100, delta=0.01)

    def test_speed_breaker_span_and_multigeometry(self):
        pts = [(73.0, 18.0, "SpeedBreakerTest")]
        # 6 breakers x 0.40m + 5 gaps x 0.60m = 2.40 + 3.00 = 5.40m
        kml, boq = temp2.build_polygons(
            pts, offline=True, speed_breaker="PTBM 10mm",
            speed_quantity=0.40, speed_gap=0.60, speed_count=6,
            log=lambda *a: None
        )
        ET.fromstring(kml)
        self.assertEqual(len(boq), 1)
        self.assertEqual(boq[0]["Mark length (m)"], 5.40)
        self.assertIn("PTBM 10mm", boq[0]["Speed Breaker"])
        self.assertEqual(kml.count("<Polygon>"), 6)
        self.assertIn("<MultiGeometry>", kml)


class TestRoadwidthModule(unittest.TestCase):
    def test_global_px_monotonic(self):
        import roadwidth
        x1, _ = roadwidth._global_px(18.0, 73.0)
        x2, _ = roadwidth._global_px(18.0, 73.001)
        self.assertGreater(x2, x1)             # east increases pixel x


# ==================================================================
# Image-marker workflow: central mapping, KMZ output, clustering, BOQ
# (all network-free)
# ==================================================================
import io as _io
import zipfile as _zipfile
import config
import temp2updated


class TestConfigMapping(unittest.TestCase):
    def test_resolve_known_and_normalisation(self):
        # bare code, lowercase, decorated name, trailing punctuation all resolve
        self.assertEqual(config.item_name("A"), "CAP PTBM 20mm X 5")
        self.assertEqual(config.item_name("a"), "CAP PTBM 20mm X 5")
        self.assertEqual(config.item_name("A - stop line"), "CAP PTBM 20mm X 5")
        self.assertEqual(config.item_name("B."), "CAP PTBM 15mm X 5")

    def test_unknown_code(self):
        self.assertIsNone(config.resolve("ZZZ"))
        self.assertFalse(config.is_known("ZZZ"))
        # item_name falls back to the original name for unknown codes
        self.assertEqual(config.item_name("ZZZ"), "ZZZ")
        self.assertEqual(config.unit_cost("ZZZ"), 0.0)

    def test_resolve_by_full_item_name(self):
        # A finalised KML/KMZ labels placemarks by the full item name (not the
        # code) - it must still resolve back to the same entry (and its cost).
        self.assertIs(config.resolve("CAP PTBM 20mm X 5"), config.resolve("A"))
        self.assertEqual(config.item_name("cap ptbm 20mm x 5"), "CAP PTBM 20mm X 5")
        self.assertEqual(config.unit_cost("CAP PTBM 20mm X 5"), config.unit_cost("A"))

    def test_image_path_exists(self):
        # every mapped image should actually be present on disk
        for code in config.PLACEHOLDER_MAP:
            self.assertIsNotNone(config.image_path(code),
                                 f"image for code {code} missing on disk")


class TestImageKmz(unittest.TestCase):
    PTS = [(73.7680, 18.4291, "A"), (73.7681, 18.4292, "A"),
           (73.7682, 18.4293, "B"), (73.7683, 18.4294, "Z")]

    def test_kmz_structure(self):
        kmz, summary = temp2updated.build_image_kmz(self.PTS, offline=True, log=lambda *a: None)
        z = _zipfile.ZipFile(_io.BytesIO(kmz))
        names = z.namelist()
        self.assertIn("doc.kml", names)
        doc = z.read("doc.kml").decode("utf-8")
        # image markers use GroundOverlay with gx:LatLonQuad for road alignment (no extra Point icon)
        self.assertIn("<GroundOverlay>", doc)
        self.assertIn("<gx:LatLonQuad>", doc)
        # unknown code Z produces a Point placemark; known codes do NOT
        self.assertIn("<Point>", doc)
        # A and B share one image -> exactly one embedded file under files/
        embedded = [n for n in names if n.startswith("files/")]
        self.assertEqual(len(embedded), 1)
        # the full item name (not the code) is used as the label
        self.assertIn("CAP PTBM 20mm X 5", doc)

    def test_unknown_code_flagged(self):
        kmz, summary = temp2updated.build_image_kmz(self.PTS, offline=True, log=lambda *a: None)
        unknown = [s for s in summary if not s["known"]]
        self.assertEqual(len(unknown), 1)
        self.assertEqual(unknown[0]["code"], "Z")

    def test_kmz_roundtrips_into_boq(self):
        # The Step-1 KMZ (placemarks labelled by full item name) must parse back
        # to points and produce the same grouped BOQ as the original codes.
        codes = [(73.7680, 18.4291, "A"), (73.7681, 18.4292, "A"),
                 (73.7682, 18.4293, "B")]
        kmz, _ = temp2updated.build_image_kmz(codes, offline=True, log=lambda *a: None)
        pts = temp2updated.parse_kml_or_kmz_bytes(kmz)      # accepts KMZ bytes
        self.assertEqual(len(pts), 3)
        grouped = temp2updated.build_grouped_boq(pts)
        rows = {r["Item"]: r["Quantity"] for r in list(grouped.values())[0]}
        self.assertEqual(rows["CAP PTBM 20mm X 5"], 2)
        self.assertEqual(rows["CAP PTBM 15mm X 5"], 1)


class TestPatchClustering(unittest.TestCase):
    def test_two_distant_patches(self):
        pts = [(73.7680, 18.4291, "A"), (73.7681, 18.4292, "B"),   # Pune-ish
               (72.8311, 21.1702, "A"), (72.8312, 21.1703, "B")]   # Surat-ish
        labels = temp2updated.cluster_patches(pts)
        self.assertEqual(labels[0], labels[1])       # near points share a patch
        self.assertEqual(labels[2], labels[3])
        self.assertNotEqual(labels[0], labels[2])    # far points split
        self.assertEqual(len(set(labels)), 2)

    def test_single_area_one_patch(self):
        pts = [(73.7680, 18.4291, "A"), (73.7681, 18.4292, "B")]
        self.assertEqual(len(set(temp2updated.cluster_patches(pts))), 1)


class TestGroupedBOQ(unittest.TestCase):
    def test_grouping_and_quantities(self):
        pts = [(73.7680, 18.4291, "A"), (73.7681, 18.4292, "A"),
               (73.7682, 18.4293, "A"), (73.7683, 18.4294, "B")]
        grouped = temp2updated.build_grouped_boq(pts)
        self.assertEqual(len(grouped), 1)            # one patch
        rows = list(grouped.values())[0]
        by_item = {r["Item"]: r for r in rows}
        self.assertEqual(by_item["CAP PTBM 20mm X 5"]["Quantity"], 3)
        self.assertEqual(by_item["CAP PTBM 15mm X 5"]["Quantity"], 1)
        # cost columns present and Total = Qty x Unit Cost
        for r in rows:
            self.assertEqual(r["Total Cost"], round(r["Quantity"] * r["Unit Cost"], 2))

    def test_separate_sheets_per_patch(self):
        pts = [(73.7680, 18.4291, "A"),
               (72.8311, 21.1702, "B")]              # far apart -> 2 patches
        grouped = temp2updated.build_grouped_boq(pts)
        self.assertEqual(len(grouped), 2)
        buf = _io.BytesIO()
        temp2updated.export_grouped_excel(grouped, buf)
        buf.seek(0)
        import openpyxl
        wb = openpyxl.load_workbook(buf)
        self.assertEqual(len(wb.sheetnames), 2)      # one worksheet per patch
        # each sheet has the cost-ready header
        for ws in wb.worksheets:
            self.assertEqual([ws.cell(1, c).value for c in range(1, 7)],
                             ["Item", "Zone", "Quantity", "Strips/Item", "Unit Rate (Rs/m2)", "Total Cost"])


# ==================================================================
# NEW FEATURES: paint price by zone, image orientation, corridor tabs,
# road width in Excel  (all network-free)
# ==================================================================

class TestZonePrices(unittest.TestCase):
    def setUp(self):
        config.reset_zone_rates()
    def tearDown(self):
        config.reset_zone_rates()

    def test_defaults_and_override(self):
        self.assertEqual(config.get_zone_rate("South"), 2500.0)
        config.set_zone_rate("north", 3000)          # case-insensitive zone
        self.assertEqual(config.get_zone_rate("North"), 3000.0)
        # override survives, others stay default
        self.assertEqual(config.get_zone_rate("East"), 2400.0)

    def test_reset_restores_defaults(self):
        config.set_zone_rate("South", 9999)
        config.reset_zone_rates()
        self.assertEqual(config.get_zone_rate("South"), 2500.0)

    def test_unknown_zone_falls_back(self):
        self.assertEqual(config.get_zone_rate("Atlantis"), config.get_zone_rate(config.DEFAULT_ZONE))
        self.assertEqual(config.get_zone_rate(""), config.get_zone_rate(config.DEFAULT_ZONE))
        with self.assertRaises(ValueError):
            config.set_zone_rate("Atlantis", 100)

    def test_invalid_prices_rejected(self):
        for bad in (-1, "abc", float("nan"), float("inf")):
            with self.assertRaises(ValueError):
                config.validate_zone_rate(bad)
        # zero is allowed (free / not-yet-priced)
        self.assertEqual(config.validate_zone_rate(0), 0.0)

    def test_bulk_set_is_atomic_and_skips_blanks(self):
        before = config.get_all_zone_rates()
        with self.assertRaises(ValueError):
            config.set_zone_rates({"South": 2600, "North": -5})   # one bad -> none applied
        self.assertEqual(config.get_all_zone_rates(), before)
        config.set_zone_rates({"East": "2900", "West": "", "Central": None})
        self.assertEqual(config.get_zone_rate("East"), 2900.0)
        self.assertEqual(config.get_zone_rate("West"), before["West"])   # blank -> unchanged

    def test_zone_rate_flows_into_boq(self):
        config.set_zone_rate("North", 1234)
        pts = [(73.0, 18.0, "x"), (73.0009, 18.0009, "A"), (73.001, 18.001, "x")]
        grouped = temp2updated.build_corridor_boq(pts, zone="North", offline=True)
        row = list(grouped.values())[0][0]
        self.assertEqual(row["Unit Cost"], 1234.0)


class TestCorridorSplit(unittest.TestCase):
    @staticmethod
    def _mk(seq):
        return [(73.0 + i * 1e-4, 18.0 + i * 1e-4, n) for i, n in enumerate(seq)]

    def _names(self, seq):
        cs = temp2updated.split_into_corridors(self._mk(seq))
        return [[nm for _, (_, _, nm) in mem] for _, mem in cs]

    def test_case1_normal(self):
        self.assertEqual(self._names(["x", "p1", "p2", "p3", "x"]),
                         [["p1", "p2", "p3"]])

    def test_case2_multiple(self):
        self.assertEqual(self._names(["x", "p1", "p2", "x", "p3", "p4", "p5", "x", "p6", "x"]),
                         [["p1", "p2"], ["p3", "p4", "p5"], ["p6"]])

    def test_case3_consecutive_x_no_empty_tab(self):
        self.assertEqual(self._names(["x", "p1", "p2", "x", "x", "p3", "p4", "x"]),
                         [["p1", "p2"], ["p3", "p4"]])

    def test_case4_no_leading_x_keeps_points(self):
        self.assertEqual(self._names(["p1", "p2", "x", "p3"]),
                         [["p1", "p2"], ["p3"]])

    def test_case5_x_at_end(self):
        self.assertEqual(self._names(["x", "p1", "p2", "p3", "x"]),
                         [["p1", "p2", "p3"]])

    def test_case6_case_and_whitespace(self):
        self.assertEqual(self._names(["x", "p1", " X ", "p2", "x"]),
                         [["p1"], ["p2"]])

    def test_labels_are_corridor_numbers(self):
        cs = temp2updated.split_into_corridors(self._mk(["x", "p1", "x", "p2", "x"]))
        self.assertEqual([lab for lab, _ in cs], ["Corridor 1", "Corridor 2"])

    def test_delimiter_detection(self):
        self.assertTrue(temp2updated.has_corridor_delimiters(self._mk(["x", "A"])))
        self.assertFalse(temp2updated.has_corridor_delimiters(self._mk(["A", "B"])))
        # real codes are never mistaken for delimiters
        self.assertFalse(temp2updated._is_corridor_delimiter("A"))
        self.assertTrue(temp2updated._is_corridor_delimiter("X - end of stretch"))


class TestCorridorBOQ(unittest.TestCase):
    def setUp(self):
        config.reset_zone_rates()
    def tearDown(self):
        config.reset_zone_rates()

    PTS = [(73.0, 18.0, "x"),
           (73.0009, 18.0009, "A"), (73.0010, 18.0010, "A"), (73.0011, 18.0011, "B"),
           (73.0012, 18.0012, "x"),
           (73.0013, 18.0013, "C"),
           (73.0014, 18.0014, "x")]

    def test_grouping_into_corridor_tabs(self):
        grouped = temp2updated.build_corridor_boq(self.PTS, zone="South", offline=True)
        self.assertEqual(list(grouped.keys()), ["Corridor 1", "Corridor 2"])
        c1 = {r["Item"]: r for r in grouped["Corridor 1"]}
        self.assertEqual(c1["CAP PTBM 20mm X 5"]["Quantity"], 2)
        self.assertEqual(c1["CAP PTBM 15mm X 5"]["Quantity"], 1)

    def test_road_width_column_present(self):
        grouped = temp2updated.build_corridor_boq(self.PTS, zone="South", offline=True)
        for rows in grouped.values():
            for r in rows:
                self.assertIn("Road Width (m)", r)
                # offline -> OSM class-default width (a number), never a crash
                self.assertIsInstance(r["Road Width (m)"], (int, float))
                self.assertGreater(r["Road Width (m)"], 0)

    def test_excel_one_sheet_per_corridor(self):
        import openpyxl
        grouped = temp2updated.build_corridor_boq(self.PTS, zone="South", offline=True)
        buf = _io.BytesIO()
        temp2updated.export_corridor_excel(grouped, buf)
        buf.seek(0)
        wb = openpyxl.load_workbook(buf)
        self.assertEqual(wb.sheetnames, ["Corridor 1", "Corridor 2"])
        hdr = [wb.worksheets[0].cell(1, c).value for c in range(1, 8)]
        self.assertEqual(hdr, ["Item", "Zone", "Quantity", "Strips/Item",
                               "Unit Rate (Rs/m2)", "Road Width (m)", "Total Cost"])
        # last data row of sheet 1 is the TOTAL row
        ws = wb["Corridor 1"]
        last = ws.max_row
        self.assertEqual(ws.cell(last, 1).value, "TOTAL")

    def test_width_measurement_failure_is_graceful(self):
        # _road_width_for_point must never raise, even with junk elements.
        def bad_get(lat, lon):
            raise RuntimeError("network down")
        w, src = temp2updated._road_width_for_point(bad_get, 73.0, 18.0, "A", 1)
        self.assertIsNone(w)
        self.assertEqual(src, "error")


class TestImageOrientation(unittest.TestCase):
    def test_canonical_along_folds_to_upper_hemisphere(self):
        import numpy as np, math
        for e, n in [(1, 0), (-1, 0), (0, 1), (0, -1), (0.7, 0.7), (-0.7, -0.7)]:
            tan = temp2updated._canonical_along(np.array([e, n], float))
            bearing = math.degrees(math.atan2(tan[0], tan[1])) % 360
            self.assertLess(bearing, 180.0 + 1e-6)

    def test_orientation_independent_of_node_order(self):
        # A road digitised forwards vs backwards must yield the same overlay axis.
        import numpy as np
        fwd = temp2updated._canonical_along(np.array([0.5, 0.8]))
        rev = temp2updated._canonical_along(np.array([-0.5, -0.8]))
        self.assertTrue(np.allclose(fwd, rev))

    def test_overlay_up_aligns_with_road_bearing(self):
        # The image "up" edge of the straight strip must point along the road
        # tangent for N/S, E/W and diagonal roads (curve-local direction).
        import numpy as np, math
        frame = temp2updated.LocalFrame(18.0, 73.0)
        for bearing in (0, 45, 90, 135):
            br = math.radians(bearing)
            tan = temp2updated._canonical_along(np.array([math.sin(br), math.cos(br)]))
            _, corners = temp2updated._straight_strip(frame, np.array([0.0, 0.0]),
                                                      tan, 3.0, 3.0, 4.0)
            # build_image_kmz maps corners as bl=1, br=2, tr=3, tl=0
            bot_mid = (corners[1] + corners[2]) / 2
            top_mid = (corners[0] + corners[3]) / 2
            up = top_mid - bot_mid
            up_bearing = math.degrees(math.atan2(up[0], up[1])) % 360
            expected = math.degrees(math.atan2(tan[0], tan[1])) % 360
            self.assertAlmostEqual(up_bearing, expected, delta=1.0)


class TestImageKmzOrientationE2E(unittest.TestCase):
    def test_kmz_has_bearing_in_summary(self):
        pts = [(73.0, 18.0, "A"), (73.0009, 18.0009, "B")]
        _, summary = temp2updated.build_image_kmz(pts, offline=True, log=lambda *a: None)
        # offline still records a (default east-west) bearing for known codes
        for s in summary:
            if s["known"]:
                self.assertIsNotNone(s["bearing"])
                self.assertTrue(0 <= s["bearing"] < 360)


if __name__ == "__main__":
    unittest.main(verbosity=2)
